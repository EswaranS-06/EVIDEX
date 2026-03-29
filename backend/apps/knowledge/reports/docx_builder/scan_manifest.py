from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def apply_page_border(section):
    sectPr = section._sectPr

    pgBorders = OxmlElement('w:pgBorders')

    for side in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '12')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        pgBorders.append(border)

    sectPr.append(pgBorders)


def draw_scan_manifest(doc, data):

    def s(x):
        return "" if x is None else str(x)

    section = doc.sections[-1]
    apply_page_border(section)

    doc.add_paragraph("")
    doc.add_paragraph("")

    # 🔥 SINGLE SOURCE WIDTH (IMPORTANT)
    table_width = 6.2  # keep SAME everywhere

    # -------------------------
    # BLUE HEADER (PERFECT FIX)
    # -------------------------

    title_table = doc.add_table(rows=1, cols=1)
    title_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    title_table.autofit = False

    tbl = title_table._element
    tblPr = tbl.tblPr

    # FIXED layout
    tblLayout = OxmlElement('w:tblLayout')
    tblLayout.set(qn('w:type'), 'fixed')
    tblPr.append(tblLayout)

    # EXACT WIDTH
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), str(int(table_width * 1440)))
    tblW.set(qn('w:type'), 'dxa')
    tblPr.append(tblW)

    cell = title_table.rows[0].cells[0]
    cell.width = Inches(table_width)

    # 🔥 HEIGHT CONTROL (THIS WAS MISSING)
    row = title_table.rows[0]
    row.height = Inches(0.5)

    tcPr = cell._element.get_or_add_tcPr()

    # Background
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "0070C0")
    tcPr.append(shd)

    # 🔥 PERFECT BORDER (MATCH TABLE)
    tcBorders = OxmlElement('w:tcBorders')
    for side in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '6')  # thinner → matches table
        border.set(qn('w:color'), '000000')
        tcBorders.append(border)
    tcPr.append(tcBorders)

    # TEXT
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)

    run = p.add_run("Scan Manifest")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(255, 255, 255)

    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # -------------------------
    # DATA
    # -------------------------

    raw_urls = data.get("target", "") or ""
    url_list = [u.strip() for u in raw_urls.split("\n") if u.strip()]
    url_count = len(url_list)

    raw_tools = data.get("tools_used", "") or ""
    tools_cleaned = ", ".join([t.strip() for t in raw_tools.split(",") if t.strip()])

    rows = [
        ("a", "Description", "Web Application Penetration Testing"),
        ("b", "Test started on", s(data.get("start_date"))),
        ("c", "Test Completed on", s(data.get("end_date"))),
        ("d", "No. of URL’s tested", str(url_count)),
        ("e", "Standard / Test Procedure reference", "OWASP TOP 10, SANS 25"),
        ("f", "Test performed at", s(data.get("test_location"))),
        ("g", "Tool used for testing", tools_cleaned),
    ]

    # -------------------------
    # TABLE (PERFECT MATCH)
    # -------------------------

    table = doc.add_table(rows=len(rows), cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # 🔥 FORCE SAME WIDTH AS HEADER
    tbl = table._element
    tblPr = tbl.tblPr

    tblLayout = OxmlElement('w:tblLayout')
    tblLayout.set(qn('w:type'), 'fixed')
    tblPr.append(tblLayout)

    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), str(int(table_width * 1440)))
    tblW.set(qn('w:type'), 'dxa')
    tblPr.append(tblW)

    # 🔥 COLUMN WIDTHS (balanced)
    col_widths = [Inches(0.4), Inches(3.2), Inches(2.6)]

    for row in table.rows:
        for i, width in enumerate(col_widths):
            row.cells[i].width = width

    # DATA FILL
    for i, (a, b, c_val) in enumerate(rows):

        row = table.rows[i]
        row.height = Inches(0.65)  # slightly taller

        cell_a, cell_b, cell_c = row.cells

        cell_a.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_b.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        p = cell_a.paragraphs[0]
        run = p.add_run(a)
        run.bold = True
        run.font.size = Pt(10)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        p = cell_b.paragraphs[0]
        run = p.add_run(b)
        run.bold = True
        run.font.size = Pt(10)

        p = cell_c.paragraphs[0]
        run = p.add_run(s(c_val))
        run.bold = True
        run.font.size = Pt(10)

    doc.add_page_break()