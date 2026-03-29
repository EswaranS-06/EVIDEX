from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION


def draw_methodology(doc, data, page_no=1, total_pages=1):

    def s(x):
        return "" if x is None else str(x)

    section = doc.sections[-1]

    # ================= PAGE BORDER =================
    sectPr = section._sectPr
    pgBorders = OxmlElement('w:pgBorders')

    for side in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '12')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        pgBorders.append(border)

    sectPr.append(pgBorders)

    # ================= TITLE =================
    p = doc.add_paragraph()
   
    run = p.add_run("2. Web Application Penetration Testing Methodology")
    run.bold = True
    run.font.size = Pt(14)
     
    

    # slight spacing
    p.paragraph_format.space_after = Pt(8)

    # ================= BULLETS =================
    bullets = [
        "Information Gathering",
        "Enumeration",
        "Scanning",
        "Exploitation",
        "Reporting",
    ]

    for b in bullets:
        p = doc.add_paragraph(f"• {b}")
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.space_after = Pt(3)

        run = p.runs[0]
        run.bold = True
        run.font.size = Pt(10)

    # ================= INTRO =================
    intro = doc.add_paragraph()

    intro.paragraph_format.left_indent = Inches(0.2)
    intro.paragraph_format.space_before = Pt(6)
    intro.paragraph_format.space_after = Pt(6)
    intro.paragraph_format.line_spacing = 1.6

    run = intro.add_run(
        "The following also gives a high-level description and process of Security Analysts "
        "methodology used for performing the Web application testing:"
    )
    run.bold = True
    run.font.size = Pt(10)

    # ================= IMAGE =================
    img_path = "apps/knowledge/reports/assets/stepss.png"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)

    run = p.add_run()
    run.add_picture(img_path, width=Inches(6.5))

    p.paragraph_format.space_after = Pt(10)

    # ================= DETAILS =================

    def add_step(title, content):
        # Title
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(6)

        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(10)

        # Content
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.line_spacing = 1.4

        run = p.add_run(content)
        run.font.size = Pt(10)

    add_step(
        "1. Planning and Reconnaissance",
        "In this initial phase, the scope and objectives of the penetration test are defined. "
        "The tester gathers relevant information about the target system through documentation review "
        "and publicly available sources to understand the environment."
    )

    add_step(
        "2. Scanning",
        "The tester uses automated and manual tools to identify vulnerabilities such as open ports, "
        "weak credentials, and misconfigurations. This phase helps determine potential entry points for exploitation."
    )

    add_step(
        "3. Gaining Access",
        "The identified vulnerabilities are exploited to gain unauthorized access. Techniques may include SQL injection, "
        "password attacks, or social engineering, depending on the defined scope."
    )

    add_step(
        "4. Maintaining Access (Optional)",
        "If permitted, the tester attempts to establish persistence within the compromised system. "
        "This phase evaluates lateral movement, privilege escalation, and the overall impact of sustained unauthorized access."
    )

    add_step(
        "5. Reporting",
        "After completing the assessment, a detailed report is prepared outlining the vulnerabilities identified, "
        "their risk level, and potential business impact. The report also provides clear and prioritized remediation recommendations."
    )

    doc.add_section(WD_SECTION.NEW_PAGE)