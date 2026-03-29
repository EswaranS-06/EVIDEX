from apps.knowledge.models import ReportFinding

def draw_conclusion(doc, data, report_id):

    doc.add_page_break()

    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    # =============================
    # CHECK CONDITION
    # =============================
    findings = ReportFinding.objects.filter(report_id=report_id)

    # Check if ALL are pending
    all_pending = all((f.status or "Pending") == "Pending" for f in findings)

    # =============================
    # TITLE
    # =============================
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)

    run = p.add_run("5. Conclusion")
    run.bold = True
    run.font.size = Pt(14)

    # =============================
    # DYNAMIC TEXT
    # =============================
    if all_pending:
        conclusion_text = (
            f"As part of the security assessment conducted for {data['application_name']}, multiple vulnerabilities "
            "were identified that may impact the confidentiality, integrity, and availability of the application. "
            "These findings indicate gaps in the current security controls and highlight areas requiring immediate attention. "
            "At the time of this report, the identified vulnerabilities remain unremediated and may expose the application "
            "to potential exploitation if left unaddressed.\n\n"
            
            "It is strongly recommended that the organization prioritize the remediation of these findings based on their "
            "severity and potential business impact. Implementing appropriate security controls, secure coding practices, "
            "and configuration hardening measures will significantly reduce the attack surface. Additionally, a formal "
            "validation or retesting exercise should be conducted post-remediation to ensure that the identified risks "
            "have been effectively mitigated.\n\n"

            "Furthermore, adopting a continuous security approach, including regular vulnerability assessments, penetration "
            "testing, and secure development lifecycle (SDLC) practices, will help in proactively identifying and addressing "
            "emerging threats, thereby strengthening the overall security posture of the application."
        )
    else:
        conclusion_text = (
            f"As part of the security assessment conducted for {data['application_name']}, multiple vulnerabilities were "
            "identified and subsequently addressed through remediation efforts. The corrective actions undertaken demonstrate "
            "a positive commitment towards improving the security posture of the application. Based on the current assessment, "
            "the remediated controls have reduced the overall risk exposure associated with the identified vulnerabilities.\n\n"

            "While the implemented fixes significantly enhance the security baseline, it is important to recognize that the "
            "threat landscape is continuously evolving. Therefore, reliance on point-in-time remediation alone is insufficient. "
            "We recommend conducting periodic security assessments, including vulnerability scanning and penetration testing, "
            "to ensure sustained protection against emerging threats.\n\n"

            "In addition, integrating security best practices into the development lifecycle, such as secure coding standards, "
            "regular code reviews, and automated security testing, will further strengthen resilience. A follow-up validation "
            "exercise is also recommended to confirm the effectiveness of the applied remediations and to ensure that no "
            "residual risks remain within the application environment."
        )
    # =============================
    # PARAGRAPH
    # =============================
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.right_indent = Inches(0.3)
    p.paragraph_format.line_spacing = 1.4
    p.paragraph_format.space_after = Pt(20)

    run = p.add_run(conclusion_text)
    run.font.size = Pt(10)

    # =============================
    # END LINE
    # =============================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)

    run = p.add_run("----END OF THE DOCUMENT----")
    run.bold = True
    run.font.size = Pt(10)