from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from django.conf import settings
import os


def remove_table_cell_margins(table):
    tbl = table._element
    tblPr = tbl.tblPr

    tblCellMar = OxmlElement('w:tblCellMar')

    for m in ["top", "left", "bottom", "right"]:
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), "0")
        node.set(qn("w:type"), "dxa")
        tblCellMar.append(node)

    tblPr.append(tblCellMar)


def draw_cover(doc, d):

    def s(x):
        return "" if x is None else str(x)

    enterprise = d.get("enterprise")
    pt_date = d.get("pt_date")
    conducted_by = d.get("application_name")

    version = s(d.get("version"))
    assessee = s(d.get("assessee"))
    assessor = s(d.get("assessor"))
    reviewed_by = s(d.get("reviewed_by"))
    approved_by = s(d.get("approved_by"))

    section = doc.sections[0]

    # -------------------------
    # PAGE MARGINS
    # -------------------------
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    # -------------------------
    # PAGE BORDER
    # -------------------------
    sectPr = section._sectPr
    pgBorders = OxmlElement('w:pgBorders')

    for side in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '12')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        pgBorders.append(border)

    sectPr.append(pgBorders)

    # -------------------------
    # HEADER
    # -------------------------
    header = section.header
    header_table = header.add_table(rows=1, cols=2, width=Inches(7.27))
    header_table.autofit = False
    remove_table_cell_margins(header_table)
    
    client_name_spacing = 7
    
    header_table.rows[0].cells[0].text = " "*client_name_spacing + enterprise
    header_table.rows[0].cells[1].text = "Penetration Testing Report"
    header_table.rows[0].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    # -------------------------
    # FOOTER
    # -------------------------
    footer = section.footer
    footer_table = footer.add_table(rows=1, cols=3, width=Inches(7.27))
    footer_table.autofit = False
    remove_table_cell_margins(footer_table)

    footer_table.rows[0].cells[0].text = "Confidential"
    footer_table.rows[0].cells[1].text = f"V {version}"
    footer_table.rows[0].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = footer_table.rows[0].cells[2].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("Page ")

    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')

    instrText = OxmlElement('w:instrText')
    instrText.text = "PAGE"

    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')

    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)

    # =========================
    # LOGOS (BIG + CLEAN)
    # =========================
    logo_table = doc.add_table(rows=1, cols=2)
    logo_table.autofit = False
    remove_table_cell_margins(logo_table)

    left_path = os.path.join(
        settings.BASE_DIR,
        "apps", "knowledge", "reports", "assets", "evidex.png"
    )

    right_path = os.path.join(
        settings.BASE_DIR,
        "apps", "knowledge", "reports", "assets", "evidex_client.png"
    )

    # LEFT LOGO
    p_left = logo_table.cell(0, 0).paragraphs[0]
    p_left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p_left.add_run()
    run.add_picture(left_path, width=Inches(2.2))   # 🔥 BIG LOGO

    # RIGHT LOGO
    p_right = logo_table.cell(0, 1).paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p_right.add_run()
    run.add_picture(right_path, width=Inches(2.2))  # 🔥 BIG LOGO

    # =========================
    # PUSH CONTENT DOWN
    # =========================
    for _ in range(1):   # 🔥 controls vertical shift
        doc.add_paragraph("")

    # =========================
    # TITLE
    # =========================
    p = doc.add_paragraph("PENETRATION TESTING REPORT")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.bold = True
    run.font.size = Pt(14)

    doc.add_paragraph("FOR").alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph(enterprise)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True

    doc.add_paragraph("")

    doc.add_paragraph(f"PT Conducted on {pt_date}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Conducted by {conducted_by}").alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("")
    doc.add_paragraph("")

    # =========================
    # TABLE
    # =========================
    table = doc.add_table(rows=5, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER

    col_widths = [1.4, 2.4, 1.2, 0.9]

    for i, width in enumerate(col_widths):
        table.columns[i].width = Inches(width)

    rows = [
        ("Document Type", "Penetration Testing Report", "Version", version),
        ("Assessee", assessee, "Signature", ""),
        ("Assessor", assessor, "Signature", ""),
        ("Reviewer", reviewed_by, "Signature", ""),
        ("Approved by", approved_by, "Signature", ""),
    ]

    for i, r in enumerate(rows):
        row = table.rows[i]

        for j in range(4):
            cell = row.cells[j]
            p = cell.paragraphs[0]

            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)

            run = p.add_run(s(r[j]))
            run.font.size = Pt(9)

    doc.add_page_break()        