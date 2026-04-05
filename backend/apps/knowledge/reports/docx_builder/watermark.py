from docx.oxml import parse_xml


def draw_watermark_docx(doc, data):

    wm = data.get("watermark_data", {})

    username = wm.get("username", "")
    ip = wm.get("ip", "")
    timestamp = wm.get("timestamp", "")

    watermark_text = f"{username} | {ip} | {timestamp}"

    section = doc.sections[0]
    header = section.header

    # ✅ USE EXISTING PARAGRAPH (NO EXTRA SPACE)
    if header.paragraphs:
        p = header.paragraphs[0]
    else:
        p = header.add_paragraph()

    watermark_xml = f"""
    <w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
         xmlns:v="urn:schemas-microsoft-com:vml"
         xmlns:o="urn:schemas-microsoft-com:office:office">
        <w:pict>
            <v:shape id="PowerPlusWaterMarkObject"
                o:spid="_x0000_s1025"
                type="#_x0000_t136"
                style="position:absolute;
                       width:480pt;
                       height:30pt;
                       rotation:315;
                       z-index:-251654144;
                       mso-position-horizontal:center;
                       mso-position-horizontal-relative:page;
                       mso-position-vertical:center;
                       mso-position-vertical-relative:page"
                fillcolor="#D9D9D9"
                stroked="f">

                <v:textpath
                    style="font-family:Calibri;
                           font-size:18pt;
                           font-weight:bold"
                    string="{watermark_text}"/>
            </v:shape>
        </w:pict>
    </w:r>
    """

    p._element.append(parse_xml(watermark_xml))