from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

from docx.enum.section import WD_SECTION


def draw_executive_summary(doc, data, page_no=1, total_pages=1):

    def s(x):
        return "" if x is None else str(x)

    section = doc.sections[-1]

    # -------------------------
    # PAGE BORDER
    # -------------------------
    sectPr = section._sectPr
    pgBorders = OxmlElement('w:pgBorders')

    for side in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '12')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        pgBorders.append(border)

    sectPr.append(pgBorders)

    # -------------------------
    # TITLE
    # -------------------------
    p = doc.add_paragraph()
    
    run = p.add_run("1. Executive Summary")
    run.bold = True
    run.font.size = Pt(14)

    doc.add_paragraph("")

    # -------------------------
    # 1.1 OVERVIEW
    # -------------------------
    p = doc.add_paragraph()
    run = p.add_run("1.1 Overview")
    run.bold = True
    run.font.size = Pt(11)

    body = doc.add_paragraph()
    body.paragraph_format.left_indent = Inches(0.2)
    body.paragraph_format.right_indent = Inches(0.1)
    body.paragraph_format.line_spacing = 1.4
    

    text = (
        "A security assessment was conducted to evaluate the effectiveness of existing controls and to identify "
        "vulnerabilities that may impact the confidentiality, integrity, and availability of the assessed environment.\n\n"
        "The engagement was performed using industry-recognized testing methodologies and simulated real-world "
        "attack scenarios to assess potential exposure to security threats. The objective was to identify exploitable "
        "weaknesses that could result in unauthorized access, data compromise, privilege escalation, or service disruption.\n\n"
        "The assessment identified findings across multiple severity levels. Each observation has been risk-rated "
        "based on standardized classification criteria and includes detailed technical analysis, impact evaluation, "
        "and recommended remediation measures.\n\n"
        "Timely remediation of identified high-risk vulnerabilities is recommended to reduce overall exposure and "
        "strengthen the organization's security posture."
    )

    run = body.add_run(text)
    run.font.size = Pt(10)

    doc.add_paragraph("")

    # -------------------------
    # 1.2 RISK MODEL
    # -------------------------
    p = doc.add_paragraph()
    run = p.add_run("1.2 Risk Model")
    run.bold = True
    run.font.size = Pt(11)

    risk_para = doc.add_paragraph()
    risk_para.paragraph_format.left_indent = Inches(0.2)
    risk_para.paragraph_format.right_indent = Inches(0.1)
    risk_para.paragraph_format.line_spacing = 1.4
    run = risk_para.add_run(
        f"Throughout this document, {s(data.get('created_by'))} has categorized the risk ratings for discovered "
        f"vulnerabilities based on global standard risk definitions."
    )
    run.font.size = Pt(10)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)

    # -------------------------
    # TABLE (FIXED - NO SPLIT)
    # -------------------------
    table = doc.add_table(rows=6, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # 🔥 Slightly tighter width (so it fits like PDF)
    col_widths = [Inches(0.8), Inches(1.3), Inches(1.2), Inches(3.8)]

    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

    # 🔥 HEADER
    header_row = table.rows[0]
    header_row.height = Inches(0.4)
    header_row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    
    headers = ["Priority\nLevel", "Severity\nScale", "CVSS\nScore", "Description of Vulnerability"]

    for j, text in enumerate(headers):
        cell = table.rows[0].cells[j]

        # ✅ vertical center
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        p = cell.paragraphs[0]

        # ✅ THIS gives internal padding
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.2

        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)

        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "00B0F0")
        tcPr.append(shd)
    
    # -------------------------
    # DATA
    # -------------------------
    data_rows = [
        ("P1", "Critical", "9.0 – 10.0",
        "The exposure may be exploited resulting in bad outcomes such as unauthorized privilege escalation, data access, downtime, or compromise of data."),
        ("P2", "High", "7.0 – 8.9",
        "These issues identify conditions that could directly result in the compromise or unauthorized access of a network, system, application, or sensitive information."),
        ("P3", "Medium", "4.0 – 6.9",
        "These issues identify conditions that do not immediately or directly result in compromise  or unauthorized access of a network, system, application, or sensitive information, but do provide a capability or information that could in combination with others’ capabilities or information result in the compromise unauthorized access of a network application or information."),
        ("P4", "Low", "0.1 – 3.9",
        "These issues identify conditions that do not immediately or directly result in the compromise of a network, system, application, or information but do provide information that could be used in combination with others’ information that could be used in combination with other's information access to a network system,application,or information."),
        ("P5", "Informational", "0",
        "Issues that leaking very basic information which might lead to information disclosure"),
    ]

    colors = ["C00000", "FF0000", "FFC000", "0070C0", "8EA9DB"]

    for i, row_data in enumerate(data_rows, start=1):

        row = table.rows[i]
        
        tr = row._element
        trPr = tr.get_or_add_trPr()

        cantSplit = OxmlElement('w:cantSplit')
        trPr.append(cantSplit)

        row.height = Inches(0.55)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

        for j, val in enumerate(row_data):
            cell = row.cells[j]

            # ✅ THIS IS THE FIX
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            p = cell.paragraphs[0]

            if j < 3:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            run = p.add_run(val)
            run.font.size = Pt(9)

            if j == 1:
                tcPr = cell._element.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), colors[i-1])
                tcPr.append(shd)

                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)   
                    
    doc.add_section(WD_SECTION.NEW_PAGE)