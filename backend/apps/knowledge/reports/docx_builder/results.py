from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from apps.knowledge.models import ReportFinding
import matplotlib.pyplot as plt
import os


# =========================
# ROW HEIGHT FIX
# =========================
def set_row_height(row, height):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), str(height))
    trHeight.set(qn('w:hRule'), 'atLeast')
    trPr.append(trHeight)


# =========================
# CHART GENERATOR
# =========================
def generate_chart_image(counts, output_path):

    labels = ["Critical", "High", "Medium", "Low", "Info"]
    values = [
        counts["CRITICAL"],
        counts["HIGH"],
        counts["MEDIUM"],
        counts["LOW"],
        counts["INFO"],
    ]

    colors = ["#C00000", "#FF0000", "#FFC000", "#0070C0", "#8EA9DB"]

    fig, ax = plt.subplots(figsize=(8, 3))

    fig.patch.set_facecolor("#2E2E2E")
    ax.set_facecolor("#2E2E2E")

    bars = ax.bar(labels, values, color=colors)

    for bar in bars:
        height = bar.get_height()
        ax.text(
            bar.get_x() + bar.get_width()/2,
            height + 0.05,
            str(int(height)),
            ha='center',
            va='bottom',
            color='white',
            fontsize=10,
            fontweight='bold'
        )

    ax.spines['bottom'].set_color('white')
    ax.spines['left'].set_color('white')
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)

    ax.tick_params(axis='x', colors='white')
    ax.tick_params(axis='y', colors='white')

    ax.set_ylim(0, max(values) + 1)

    plt.tight_layout()
    plt.savefig(output_path, dpi=300, facecolor=fig.get_facecolor())
    plt.close()


# =========================
# MAIN FUNCTION
# =========================
def draw_results(doc, data, report_id):

    SEVERITY_ORDER = {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 2, "LOW": 3, "INFO": 4}
    STATUS_ORDER = {"Pending": 0, "Patched": 1}

    findings = [
        f for f in ReportFinding.objects.filter(report_id=report_id)
        if (f.final_title and f.final_severity)
    ]

    findings.sort(
        key=lambda f: (
            STATUS_ORDER.get(f.status or "Pending", 1),
            SEVERITY_ORDER.get((f.final_severity or "").upper(), 5)
        )
    )

    # =========================
    # COUNTS
    # =========================
    counts = {"CRITICAL": 0, "HIGH": 0, "MEDIUM": 0, "LOW": 0, "INFO": 0}
    for f in findings:
        sev = (f.final_severity or "").upper()
        if sev in counts:
            counts[sev] += 1

    # =========================
    # TITLE
    # =========================
    p = doc.add_paragraph()
    
    run = p.add_run("3. Project Scope")
    run.bold = True
    run.font.size = Pt(14)

    scope = doc.add_paragraph()
    scope.paragraph_format.left_indent = Inches(0.2)

    run = scope.add_run(
        f"Formal communication from the {data['application_name']} outlined the application "
        "to be tested and the type of testing to be carried out. "
        "A RED team resource was deployed to perform this activity."
    )
    run.font.size = Pt(10)

    doc.add_paragraph("")

    # =========================
    # RESULTS TITLE
    # =========================
    p = doc.add_paragraph()
   
    run = p.add_run("4. Penetration Testing Results")
    run.bold = True
    run.font.size = Pt(14)

    doc.add_paragraph("")

    # =========================
    # SUMMARY TABLE
    # =========================
    table = doc.add_table(rows=2, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["URL", "Critical", "High", "Medium", "Low", "Information"]
    values = [
        f"https://{data['application_name']}",
        counts["CRITICAL"],
        counts["HIGH"],
        counts["MEDIUM"],
        counts["LOW"],
        counts["INFO"],
    ]

    colors = ["0070C0", "C00000", "FF0000", "FFC000", "0070C0", "8EA9DB"]

    for j in range(6):
        cell = table.rows[0].cells[j]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = p.add_run(headers[j])
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)

        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), colors[j])
        tcPr.append(shd)

        cell = table.rows[1].cells[j]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = p.add_run(str(values[j]))
        run.bold = True
        run.font.size = Pt(9)

    # ✅ FIX HEIGHT (IMPORTANT)
    set_row_height(table.rows[0], 500)
    set_row_height(table.rows[1], 500)

    doc.add_paragraph("")

    # =========================
    # CHART
    # =========================
    chart_path = "chart.png"
    generate_chart_image(counts, chart_path)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run()
    run.add_picture(chart_path, width=Inches(6.5))

    if os.path.exists(chart_path):
        os.remove(chart_path)

    doc.add_paragraph("")

    # =========================
    # FINDINGS TABLE
    # =========================
    if findings:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = ["S.No", "Vulnerability Name", "Severity", "Status"]

        for j in range(4):
            cell = table.rows[0].cells[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            run = p.add_run(headers[j])
            run.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(255, 255, 255)

            tcPr = cell._element.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "1F4E8C")
            tcPr.append(shd)

        # header height
        set_row_height(table.rows[0], 550)

    for i, f in enumerate(findings, start=1):
        row = table.add_row()
        set_row_height(row, 500)

        severity = (f.final_severity or "").upper()
        status = f.status or "Pending"

        vals = [i, f.final_title, severity, status]

        for j, val in enumerate(vals):
                cell = row.cells[j]
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

                p = cell.paragraphs[0]

                if j != 1:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

                run = p.add_run(str(val))
                run.font.size = Pt(9)

                # =========================
                # 🎨 SEVERITY COLOR
                # =========================
                if j == 2:
                    color_map = {
                        "CRITICAL": RGBColor(192, 0, 0),
                        "HIGH": RGBColor(255, 0, 0),
                        "MEDIUM": RGBColor(255, 192, 0),
                        "LOW": RGBColor(0, 112, 192),
                        "INFO": RGBColor(142, 169, 219),
                    }
                    run.bold = True
                    run.font.color.rgb = color_map.get(severity, RGBColor(0, 0, 0))

                # =========================
                # 🎨 STATUS COLOR
                # =========================
                if j == 3:
                    status_map = {
                        "Pending": RGBColor(192, 0, 0),   # red
                        "Patched": RGBColor(0, 176, 80),  # green
                    }
                    run.bold = True
                    run.font.color.rgb = status_map.get(status, RGBColor(0, 0, 0))
                    
    doc.add_page_break()